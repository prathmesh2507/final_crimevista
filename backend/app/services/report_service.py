import os
import re
import uuid
from collections import Counter
from datetime import datetime, timezone
from html import escape
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as PdfTable,
    TableStyle,
)
from sqlalchemy.orm import Session

from ..models import Crime
from .analytics_service import build_insights, build_kpis, counts
from .data_service import CRIME_COLUMNS
from .filter_service import filtered_crimes

REPORTS: list[dict] = []
PDF_RECORD_LIMIT = 300
SEVERITY_ORDER = ["Critical", "High", "Medium", "Low"]
REPORT_SECTIONS = {"executive_summary", "kpis", "charts", "hotspots", "records"}


def _reports_dir() -> Path:
    path = Path(
        os.getenv("REPORTS_DIR", Path(__file__).resolve().parents[2] / "reports")
    )
    if not path.is_absolute():
        path = Path(__file__).resolve().parents[2] / path
    path.mkdir(parents=True, exist_ok=True)
    return path


def _file_name(title: str, report_id: str, extension: str) -> str:
    safe_title = (
        re.sub(r"[^A-Za-z0-9_-]+", "_", title).strip("_")[:60] or "crimevista_report"
    )
    return f"{safe_title}_{report_id}.{extension}"


def _records_frame(rows: list[Crime]) -> pd.DataFrame:
    return pd.DataFrame(
        [{column: getattr(row, column) for column in CRIME_COLUMNS} for row in rows],
        columns=CRIME_COLUMNS,
    )


def _date_range(rows: list[Crime]) -> str:
    dates = [row.date for row in rows if row.date]
    if not dates:
        return "No dated records"
    first, last = min(dates), max(dates)
    return (
        first.isoformat()
        if first == last
        else f"{first.isoformat()} to {last.isoformat()}"
    )


def _distribution_rows(rows: list[Crime]) -> list[list[object]]:
    total = len(rows)
    result = []
    groups = [
        ("Crime type", "crime_type", None),
        ("Severity", "severity", SEVERITY_ORDER),
        ("Time of day", "time_period", None),
        ("Area", "area", None),
    ]
    for label, field, order in groups:
        for value, count in counts(rows, field, ordered_values=order):
            result.append([label, value, count, count / total if total else 0])

    months = Counter(row.date.strftime("%Y-%m") for row in rows if row.date)
    for month, count in sorted(months.items()):
        result.append(["Month", month, count, count / total if total else 0])
    return result


def _hotspot_rows(rows: list[Crime], limit: int = 20) -> list[list[object]]:
    total = len(rows)
    area_counts = counts(rows, "area", limit)
    result = []
    for rank, (area, count) in enumerate(area_counts, start=1):
        crime_types = Counter(
            row.crime_type for row in rows if row.area == area and row.crime_type
        )
        dominant_type = sorted(
            crime_types.items(), key=lambda item: (-item[1], item[0])
        )[0][0]
        result.append([rank, area, count, count / total if total else 0, dominant_type])
    return result


def _pdf_table(data, widths=None, font_size=8, repeat_rows=1):
    cell_style = ParagraphStyle(
        f"table_cell_{font_size}",
        fontName="Helvetica",
        fontSize=font_size,
        leading=font_size + 2,
    )
    wrapped = [
        [Paragraph(escape(str(value)), cell_style) for value in row] for row in data
    ]
    table = LongTable(wrapped, colWidths=widths, repeatRows=repeat_rows, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("LEADING", (0, 0), (-1, -1), font_size + 2),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DEE8")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F3F6FA")],
                ),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _draw_pdf_chrome(title: str, generated_at: datetime):
    def draw(canvas, document):
        canvas.saveState()
        page_width, page_height = landscape(letter)
        canvas.setStrokeColor(colors.HexColor("#D6DEE8"))
        canvas.line(
            document.leftMargin,
            page_height - 0.42 * inch,
            page_width - document.rightMargin,
            page_height - 0.42 * inch,
        )
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(colors.HexColor("#16324F"))
        canvas.drawString(document.leftMargin, page_height - 0.31 * inch, title[:90])
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#607184"))
        canvas.drawRightString(
            page_width - document.rightMargin,
            0.25 * inch,
            f"CrimeVista | {generated_at:%Y-%m-%d %H:%M UTC} | Page {document.page}",
        )
        canvas.restoreState()

    return draw


def _export_pdf(
    rows: list[Crime],
    path: Path,
    title: str,
    sections: set[str],
    generated_at: datetime,
):
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=27,
            textColor=colors.HexColor("#16324F"),
            alignment=0,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#16324F"),
            spaceBefore=14,
            spaceAfter=7,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportBody",
            parent=styles["BodyText"],
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#263746"),
            spaceAfter=5,
        )
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        leftMargin=0.55 * inch,
        rightMargin=0.55 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.55 * inch,
        title=title,
        author="CrimeVista",
    )
    story = [Paragraph(escape(title), styles["ReportTitle"])]
    story.append(Paragraph("Crime analysis report", styles["ReportBody"]))
    story.append(Spacer(1, 8))
    scope = [
        [
            "Generated (UTC)",
            generated_at.strftime("%Y-%m-%d %H:%M"),
            "Matching records",
            f"{len(rows):,}",
        ],
        [
            "Incident date range",
            _date_range(rows),
            "Areas represented",
            f"{len({row.area for row in rows if row.area}):,}",
        ],
    ]
    scope_table = PdfTable(
        scope, colWidths=[1.25 * inch, 3.0 * inch, 1.45 * inch, 1.3 * inch]
    )
    scope_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EEF5")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#263746")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DEE8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([scope_table, Spacer(1, 10)])

    if not rows:
        story.append(
            Paragraph(
                "No incidents matched the selected report filters.",
                styles["ReportBody"],
            )
        )

    if "executive_summary" in sections and rows:
        story.append(Paragraph("Executive summary", styles["ReportHeading"]))
        for insight in build_insights(rows)[:5]:
            story.append(
                Paragraph(
                    f"&#8226;&nbsp; {escape(insight['text'])}", styles["ReportBody"]
                )
            )

    if "kpis" in sections:
        story.append(Paragraph("Key indicators", styles["ReportHeading"]))
        kpis = [["Indicator", "Value", "Context"]]
        kpis.extend(
            [
                [item["label"], item["value"], item.get("context") or ""]
                for item in build_kpis(rows)
            ]
        )
        story.append(
            _pdf_table(kpis, [2.1 * inch, 1.5 * inch, 6.0 * inch], font_size=8)
        )

    if "charts" in sections:
        story.append(Paragraph("Distributions", styles["ReportHeading"]))
        distributions = [["Group", "Category", "Incidents", "Share"]]
        distributions.extend(
            [
                [group, value, f"{count:,}", f"{share:.1%}"]
                for group, value, count, share in _distribution_rows(rows)
            ]
        )
        story.append(
            _pdf_table(
                distributions,
                [1.8 * inch, 3.6 * inch, 1.5 * inch, 1.2 * inch],
                font_size=7.5,
            )
        )

    if "hotspots" in sections:
        story.append(Paragraph("Area hotspot ranking", styles["ReportHeading"]))
        hotspots = [["Rank", "Area", "Incidents", "Share", "Most common crime type"]]
        hotspots.extend(
            [
                [rank, area, f"{count:,}", f"{share:.1%}", crime_type]
                for rank, area, count, share, crime_type in _hotspot_rows(rows)
            ]
        )
        story.append(
            _pdf_table(
                hotspots,
                [0.65 * inch, 2.5 * inch, 1.4 * inch, 1.2 * inch, 3.0 * inch],
                font_size=8,
            )
        )

    if "records" in sections:
        if sections & {"executive_summary", "kpis", "charts", "hotspots"}:
            story.append(PageBreak())
        story.append(Paragraph("Incident records", styles["ReportHeading"]))
        if len(rows) > PDF_RECORD_LIMIT:
            story.append(
                Paragraph(
                    f"Showing the first {PDF_RECORD_LIMIT:,} of {len(rows):,} matching records. CSV exports and Excel workbooks with the Records section contain the complete record set.",
                    styles["ReportBody"],
                )
            )
        record_headers = [
            "Crime ID",
            "Date",
            "Area",
            "Crime type",
            "Severity",
            "Status",
            "Police station",
        ]
        record_rows = [record_headers]
        ordered_rows = sorted(
            rows, key=lambda row: (row.date or datetime.min.date(), str(row.crime_id))
        )
        record_rows.extend(
            [
                [
                    row.crime_id,
                    row.date.isoformat() if row.date else "",
                    row.area,
                    row.crime_type,
                    row.severity,
                    row.status or "",
                    row.police_station or "",
                ]
                for row in ordered_rows[:PDF_RECORD_LIMIT]
            ]
        )
        story.append(
            _pdf_table(
                record_rows,
                [
                    1.25 * inch,
                    0.85 * inch,
                    1.45 * inch,
                    1.5 * inch,
                    0.8 * inch,
                    1.4 * inch,
                    1.65 * inch,
                ],
                font_size=7,
                repeat_rows=1,
            )
        )

    page_callback = _draw_pdf_chrome(title, generated_at)
    doc.build(story, onFirstPage=page_callback, onLaterPages=page_callback)


def _write_excel_table(
    ws, headers: list[str], rows: list[list[object]], table_name: str
):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="16324F")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    if rows:
        table = Table(
            displayName=table_name,
            ref=f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)
    for column in ws.columns:
        values = [len(str(cell.value or "")) for cell in column[:200]]
        width = min(max(max(values, default=10) + 2, 12), 42)
        ws.column_dimensions[column[0].column_letter].width = width
    ws.auto_filter.ref = ws.dimensions


def _export_xlsx(
    rows: list[Crime],
    path: Path,
    title: str,
    sections: set[str],
    generated_at: datetime,
):
    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append([title])
    overview.merge_cells("A1:C1")
    overview["A1"].font = Font(size=18, bold=True, color="16324F")
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 32
    metadata = [
        ["Generated (UTC)", generated_at.strftime("%Y-%m-%d %H:%M")],
        ["Matching records", len(rows)],
        ["Incident date range", _date_range(rows)],
        ["Areas represented", len({row.area for row in rows if row.area})],
    ]
    overview.append([])
    for item in metadata:
        overview.append(item)
    for row in overview.iter_rows(
        min_row=3, max_row=overview.max_row, min_col=1, max_col=2
    ):
        row[0].font = Font(bold=True, color="16324F")
    row_cursor = overview.max_row + 2

    if "executive_summary" in sections:
        overview.cell(row_cursor, 1, "Executive summary")
        overview.cell(row_cursor, 1).font = Font(size=12, bold=True, color="16324F")
        row_cursor += 1
        insights = build_insights(rows) or [
            "No incidents matched the selected report filters."
        ]
        for insight in insights[:5]:
            overview.cell(
                row_cursor, 1, insight if isinstance(insight, str) else insight["text"]
            )
            overview.merge_cells(
                start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=3
            )
            overview.cell(row_cursor, 1).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            row_cursor += 1
        row_cursor += 1

    if "kpis" in sections:
        overview.cell(row_cursor, 1, "Key indicators")
        overview.cell(row_cursor, 1).font = Font(size=12, bold=True, color="16324F")
        row_cursor += 1
        kpi_rows = [
            [item["label"], item["value"], item.get("context") or ""]
            for item in build_kpis(rows)
        ]
        for values in [["Indicator", "Value", "Context"], *kpi_rows]:
            for column_index, value in enumerate(values, start=1):
                overview.cell(row_cursor, column_index, value)
            row_cursor += 1
        for cell in overview[row_cursor - len(kpi_rows) - 1]:
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.font = Font(color="FFFFFF", bold=True)

    overview.column_dimensions["A"].width = 31
    overview.column_dimensions["B"].width = 24
    overview.column_dimensions["C"].width = 54
    overview.freeze_panes = "A3"

    if "charts" in sections:
        sheet = workbook.create_sheet("Distributions")
        _write_excel_table(
            sheet,
            ["Group", "Category", "Incidents", "Share"],
            _distribution_rows(rows),
            "CrimeDistributions",
        )
        for cell in sheet["D"][1:]:
            cell.number_format = "0.0%"

    if "hotspots" in sections:
        sheet = workbook.create_sheet("Hotspots")
        _write_excel_table(
            sheet,
            ["Rank", "Area", "Incidents", "Share", "Most common crime type"],
            _hotspot_rows(rows),
            "CrimeHotspots",
        )
        for cell in sheet["D"][1:]:
            cell.number_format = "0.0%"

    if "records" in sections:
        sheet = workbook.create_sheet("Records")
        frame = _records_frame(rows)
        _write_excel_table(
            sheet,
            CRIME_COLUMNS,
            frame.values.tolist(),
            "CrimeRecords",
        )

    workbook.save(path)


def _export(
    rows: list[Crime],
    path: Path,
    file_format: str,
    title: str,
    sections: list[str] | None = None,
    generated_at: datetime | None = None,
):
    sections = set(sections or []) & REPORT_SECTIONS
    generated_at = generated_at or datetime.now(timezone.utc)
    frame = _records_frame(rows)
    if file_format == "csv":
        frame.to_csv(path, index=False)
    elif file_format == "xlsx":
        _export_xlsx(rows, path, title, sections, generated_at)
    else:
        _export_pdf(rows, path, title, sections, generated_at)


def generate_report(db: Session, request):
    rows = filtered_crimes(db, request.filters)
    report_id = f"rpt_{uuid.uuid4().hex[:12]}"
    path = _reports_dir() / _file_name(request.title, report_id, request.format)
    generated_at = datetime.now(timezone.utc)
    _export(rows, path, request.format, request.title, request.sections, generated_at)
    report = {
        "reportId": report_id,
        "title": request.title,
        "format": request.format,
        "status": "ready",
        "createdAt": generated_at.isoformat(),
        "downloadUrl": f"/api/reports/{report_id}/download",
        "preview": {
            "summary": build_insights(rows)[:4],
            "recordCount": len(rows),
            "kpis": build_kpis(rows),
        },
        "_path": str(path),
    }
    REPORTS.insert(0, report)
    return {key: value for key, value in report.items() if key != "_path"}


def list_reports():
    return [
        {key: value for key, value in report.items() if key != "_path"}
        for report in REPORTS
    ]


def find_report(report_id: str):
    return next((report for report in REPORTS if report["reportId"] == report_id), None)
